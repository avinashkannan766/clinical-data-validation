from openpyxl import load_workbook
from pathlib import Path
from copy import copy
from datetime import timedelta


# =========================================================
# FILES
# =========================================================

clean_file = Path("data/clinical_study_corrected.xlsx")
test_file = Path("data/clinical_study_test.xlsx")


if not clean_file.exists():
    raise FileNotFoundError(
        f"Clean workbook not found: {clean_file.resolve()}"
    )


# =========================================================
# START FROM CLEAN WORKBOOK
# =========================================================

workbook = load_workbook(clean_file)


# =========================================================
# HELPER FUNCTION
# =========================================================

def get_headers(sheet):
    return {
        cell.value: cell.column
        for cell in sheet[1]
    }


def find_row(sheet, column_name, value):
    headers = get_headers(sheet)
    column = headers[column_name]

    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, column).value == value:
            return row

    raise ValueError(
        f"{value} not found in column {column_name}"
    )


# =========================================================
# DEF-001
# Missing participant Age
# Expected: VAL-002
# =========================================================

participants = workbook["Participants"]
p_headers = get_headers(participants)

ct005_row = find_row(
    participants,
    "Subject_ID",
    "CT005"
)

participants.cell(
    ct005_row,
    p_headers["Age"]
).value = None


# =========================================================
# DEF-002
# Duplicate participant Subject_ID
# CT006 -> CT005
# Expected: VAL-001
# =========================================================

ct006_row = find_row(
    participants,
    "Subject_ID",
    "CT006"
)

participants.cell(
    ct006_row,
    p_headers["Subject_ID"]
).value = "CT005"


# =========================================================
# DEF-003
# Invalid visit date
# CT004 Week 4 -> invalid date text
# Expected: VAL-009
# =========================================================

visits = workbook["Visits"]
v_headers = get_headers(visits)

ct004_week4_row = None

for row in range(2, visits.max_row + 1):
    subject = visits.cell(
        row,
        v_headers["Subject_ID"]
    ).value

    visit = visits.cell(
        row,
        v_headers["Visit_Name"]
    ).value

    if subject == "CT004" and visit == "Week 4":
        ct004_week4_row = row
        break

if ct004_week4_row is None:
    raise ValueError("CT004 Week 4 visit not found.")

visits.cell(
    ct004_week4_row,
    v_headers["Visit_Date"]
).value = "2026-02-30"


# =========================================================
# DEF-004
# Visit outside allowed window
# CT007 Week 4 -> Day 40 instead of Day 28
# Expected: VAL-012
# =========================================================

participants_dates = {}

for row in range(2, participants.max_row + 1):
    subject = participants.cell(
        row,
        p_headers["Subject_ID"]
    ).value

    consent_date = participants.cell(
        row,
        p_headers["Consent_Date"]
    ).value

    if subject and consent_date:
        participants_dates[subject] = consent_date


ct007_week4_row = None

for row in range(2, visits.max_row + 1):
    subject = visits.cell(
        row,
        v_headers["Subject_ID"]
    ).value

    visit = visits.cell(
        row,
        v_headers["Visit_Name"]
    ).value

    if subject == "CT007" and visit == "Week 4":
        ct007_week4_row = row
        break

if ct007_week4_row is None:
    raise ValueError("CT007 Week 4 visit not found.")

ct007_consent = participants_dates["CT007"]

visits.cell(
    ct007_week4_row,
    v_headers["Visit_Date"]
).value = ct007_consent + timedelta(days=40)


# =========================================================
# DEF-005
# Orphan laboratory subject
# LAB005 -> CT999
# Expected: VAL-013
# =========================================================

laboratory = workbook["Laboratory"]
l_headers = get_headers(laboratory)

lab005_row = find_row(
    laboratory,
    "Lab_ID",
    "LAB005"
)

laboratory.cell(
    lab005_row,
    l_headers["Subject_ID"]
).value = "CT999"


# =========================================================
# DEF-006
# Duplicate laboratory record
# Duplicate LAB006
# Expected: VAL-028
# =========================================================

lab006_row = find_row(
    laboratory,
    "Lab_ID",
    "LAB006"
)

# Copy the entire LAB006 record and append it.
new_row = laboratory.max_row + 1

for column in range(1, laboratory.max_column + 1):
    source_cell = laboratory.cell(
        lab006_row,
        column
    )

    target_cell = laboratory.cell(
        new_row,
        column
    )

    target_cell.value = source_cell.value

    if source_cell.has_style:
        target_cell._style = copy(source_cell._style)

    if source_cell.number_format:
        target_cell.number_format = source_cell.number_format


# =========================================================
# DEF-007
# Unsupported laboratory unit
# LAB007 -> mg/L
# Expected: VAL-017
# =========================================================

lab007_row = find_row(
    laboratory,
    "Lab_ID",
    "LAB007"
)

laboratory.cell(
    lab007_row,
    l_headers["Unit"]
).value = "mg/L"


# =========================================================
# DEF-008
# Non-numeric laboratory result
# LAB008 -> NOT AVAILABLE
# Expected: VAL-016
# =========================================================

lab008_row = find_row(
    laboratory,
    "Lab_ID",
    "LAB008"
)

laboratory.cell(
    lab008_row,
    l_headers["Result_Value"]
).value = "NOT AVAILABLE"


# =========================================================
# DEF-009
# Laboratory linked to a missed visit
# CT008 Week 4 -> Missed
# Expected: VAL-020
# =========================================================

ct008_week4_row = None

for row in range(2, visits.max_row + 1):
    subject = visits.cell(
        row,
        v_headers["Subject_ID"]
    ).value

    visit = visits.cell(
        row,
        v_headers["Visit_Name"]
    ).value

    if subject == "CT008" and visit == "Week 4":
        ct008_week4_row = row
        break

if ct008_week4_row is None:
    raise ValueError("CT008 Week 4 visit not found.")

visits.cell(
    ct008_week4_row,
    v_headers["Visit_Status"]
).value = "Missed"


# =========================================================
# DEF-010
# AE onset before consent
# AE001 -> one day before CT001 consent
# Expected: VAL-025
# =========================================================

adverse_events = workbook["Adverse_Events"]
ae_headers = get_headers(adverse_events)

ae001_row = find_row(
    adverse_events,
    "AE_ID",
    "AE001"
)

ct001_consent = participants_dates["CT001"]

adverse_events.cell(
    ae001_row,
    ae_headers["Onset_Date"]
).value = ct001_consent - timedelta(days=1)


# =========================================================
# DEF-011
# Baseline visit before consent
# CT010 Baseline -> one day before consent
# Expected: VAL-012
# =========================================================

ct010_baseline_row = None

for row in range(2, visits.max_row + 1):
    subject = visits.cell(
        row,
        v_headers["Subject_ID"]
    ).value

    visit = visits.cell(
        row,
        v_headers["Visit_Name"]
    ).value

    if subject == "CT010" and visit == "Baseline":
        ct010_baseline_row = row
        break

if ct010_baseline_row is None:
    raise ValueError("CT010 Baseline visit not found.")

ct010_consent = participants_dates["CT010"]

visits.cell(
    ct010_baseline_row,
    v_headers["Visit_Date"]
).value = ct010_consent - timedelta(days=1)


# =========================================================
# DEF-012
# Orphan adverse-event subject
# AE012 -> CT999
# Expected: VAL-022
# =========================================================

ae012_row = find_row(
    adverse_events,
    "AE_ID",
    "AE012"
)

adverse_events.cell(
    ae012_row,
    ae_headers["Subject_ID"]
).value = "CT999"


# =========================================================
# SAVE
# =========================================================

workbook.save(test_file)


# =========================================================
# SUMMARY
# =========================================================

print("=" * 70)
print("REPRODUCIBLE TEST WORKBOOK CREATED")
print("=" * 70)

print(f"Source: {clean_file.resolve()}")
print(f"Output: {test_file.resolve()}")

print()
print("Defects introduced:")

print("DEF-001 | CT005 | Missing Age")
print("DEF-002 | CT006 changed to CT005 | Duplicate Subject_ID")
print("DEF-003 | CT004 Week 4 | Invalid Visit_Date")
print("DEF-004 | CT007 Week 4 | Visit outside allowed window")
print("DEF-005 | LAB005 | Orphan Subject_ID CT999")
print("DEF-006 | LAB006 | Duplicate laboratory record")
print("DEF-007 | LAB007 | Unsupported Unit mg/L")
print("DEF-008 | LAB008 | Non-numeric Result_Value")
print("DEF-009 | CT008 Week 4 | Visit marked Missed")
print("DEF-010 | AE001 | Onset before consent")
print("DEF-011 | CT010 Baseline | Date before consent")
print("DEF-012 | AE012 | Orphan Subject_ID CT999")

print()
print("The clean corrected workbook was not modified.")
print("=" * 70)